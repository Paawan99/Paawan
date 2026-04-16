"""Tailor a .docx resume to a specific job description.

Strategy:
- Open the user's source resume.docx (python-docx).
- Locate a section we treat as the "Skills" / "Core Competencies" block by
  matching common heading text (case-insensitive).
- Append any keywords from the JD that appear in profile.skills but are not
  already present in the resume's skills section.
- Save to a per-job filename under output/.

We deliberately DO NOT rewrite bullets or invent experience. Only the skills
block is modified, and only with terms the user has pre-declared in their
profile. Anything else would be making things up.

python-docx is loaded lazily so the rest of the CLI works without it
installed.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import List, Set

log = logging.getLogger(__name__)

# headings we'll treat as the skills section
_SKILLS_HEADINGS = (
    "skills",
    "core skills",
    "core competencies",
    "technical skills",
    "key skills",
    "technologies",
    "tech stack",
)


def _looks_like_heading(text: str) -> str | None:
    """Return the canonical heading if `text` matches a skills heading."""
    s = re.sub(r"[^a-z ]", "", (text or "").lower()).strip()
    for h in _SKILLS_HEADINGS:
        if s == h:
            return h
    return None


def _existing_skills_text(doc) -> str:
    """Concatenate paragraphs in the skills section so we can dedupe."""
    in_section = False
    chunks: List[str] = []
    for p in doc.paragraphs:
        text = p.text or ""
        if _looks_like_heading(text):
            if in_section:
                break  # next heading reached
            in_section = True
            continue
        # cheap heuristic for "next section started": ALL CAPS short line
        if in_section and text.strip() and text.strip().isupper() and len(text) < 40:
            break
        if in_section:
            chunks.append(text)
    return " | ".join(chunks).lower()


def _approved_additions(
    profile_skills: dict,
    matched_keywords: List[str],
    already_present: str,
) -> List[str]:
    """Keywords the JD called out that are in the profile but not in the resume."""
    flat: Set[str] = set()
    for terms in (profile_skills or {}).values():
        for t in terms or []:
            flat.add(t.strip())

    additions: List[str] = []
    for kw in matched_keywords:
        for term in flat:
            if term.lower() == kw.lower() and term.lower() not in already_present:
                additions.append(term)
                break
    # preserve order, dedupe case-insensitively
    seen, out = set(), []
    for a in additions:
        k = a.lower()
        if k not in seen:
            seen.add(k)
            out.append(a)
    return out


def tailor_docx(
    source_path: str | Path,
    output_path: str | Path,
    profile: dict,
    matched_keywords: List[str],
) -> dict:
    """Inject approved keywords into the skills section of the resume.

    Returns a small report so the CLI can show what changed.
    """
    try:
        from docx import Document  # type: ignore
    except ImportError as e:
        raise RuntimeError(
            "python-docx is required for tailor_docx. "
            "Install with: pip install python-docx"
        ) from e

    source_path = Path(source_path)
    output_path = Path(output_path)
    if not source_path.exists():
        raise FileNotFoundError(source_path)

    doc = Document(str(source_path))
    already = _existing_skills_text(doc)
    additions = _approved_additions(
        profile.get("skills", {}), matched_keywords, already
    )

    if not additions:
        doc.save(str(output_path))
        return {"output": str(output_path), "added": [], "note": "no changes"}

    # find the last paragraph in the skills section and append
    in_section = False
    last_in_section = None
    insert_index = None
    for i, p in enumerate(doc.paragraphs):
        text = p.text or ""
        if _looks_like_heading(text):
            if in_section:
                insert_index = i
                break
            in_section = True
            last_in_section = p
            continue
        if in_section and text.strip() and text.strip().isupper() and len(text) < 40:
            insert_index = i
            break
        if in_section:
            last_in_section = p

    addition_line = "Additional: " + ", ".join(additions)
    if last_in_section is not None:
        # append a run on a new paragraph right after the last skills paragraph
        new_p = last_in_section.insert_paragraph_before(addition_line)  # type: ignore[attr-defined]
        # insert_paragraph_before puts it BEFORE — we want after, so swap text
        new_p.text, last_in_section.text = last_in_section.text, addition_line
    else:
        # no skills section found — append at end as a new section
        doc.add_paragraph("Skills")
        doc.add_paragraph(addition_line)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return {
        "output": str(output_path),
        "added": additions,
        "note": f"added {len(additions)} keyword(s)",
    }
